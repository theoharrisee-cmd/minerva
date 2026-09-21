"""v0.3 features: mandate, contacts, intake, vault, planning register, comps, Outlook (mocked), Excel model, API."""
import base64, json, os, pathlib, shutil, subprocess, sys, tempfile, threading, urllib.request

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
TMP = tempfile.mkdtemp(prefix="minerva-test-")
os.environ.update(MINERVA_HOME=TMP, MINERVA_NO_OPEN="1", MINERVA_DOWNLOADS=TMP + "/dl")

from core import comps, contacts, demo_deals, docs, intake, investor, outlook, planreg, store, xlmodel  # noqa: E402
from core.appraisal import appraise  # noqa: E402
from core.config import DEFAULT_ASSUMPTIONS  # noqa: E402

PROFILE = {**investor.DEFAULT_INVESTOR, "active": True, "fund_size": 200, "equity_available": 20, "max_deal_pct_fund": 10, "min_deal": 5, "max_deal": 60,
           "regions": ["London"], "strategies": ["develop_sell"], "target_irr": 15, "exclusions": ["green_belt"]}


def _ev(site, strat=None):
    a = appraise(site, {**DEFAULT_ASSUMPTIONS, **site.get("ovr", {})}, strat or site.get("strategy", "develop_sell"))
    return {"appraisal": {k: a.get(k) for k in ("total_cost", "peak_equity", "units", "irr_levered", "equity_multiple", "npv_equity", "months_land", "incomplete")}, "strategy": strat or site.get("strategy")}


def test_region_and_fit():
    assert investor.region_of({"postcode": "LS10 1JQ"}) == "Yorkshire and the Humber"
    assert investor.region_of({"postcode": "SW19 2AB"}) == "London"
    s = {"postcode": "LS10 1JQ", "site_ha": 1, "sales_psf": 500, "asking_price": 2e6, "planning_status": "Consented", "strategy": "develop_sell"}
    f = investor.mandate_fit(s, _ev(s), PROFILE)
    assert f["hard_fail"] and any(c["name"] == "Region" and c["ok"] is False for c in f["checks"])
    s["postcode"] = "SW19 2AB"
    s["green_belt"] = True
    f = investor.mandate_fit(s, _ev(s), PROFILE)
    assert f["hard_fail"] and any(c["name"] == "Exclusions" and c["ok"] is False for c in f["checks"])
    assert investor.mandate_fit(s, _ev(s), {**PROFILE, "active": False})["active"] is False


def test_provisional_when_incomplete():
    s = {"postcode": "SW19 2AB", "strategy": "develop_sell", "planning_status": "Consented"}
    ev = {"appraisal": {"incomplete": True}, "strategy": "develop_sell"}
    f = investor.mandate_fit(s, ev, PROFILE)
    assert f["label"] == "Provisional" and (f["score"] or 0) <= 60


def test_derived_assumptions_and_search_filter():
    d = investor.derived_assumptions({**PROFILE, "senior_rate": 9.0, "senior_ltc": 55.0, "cost_of_equity": 14.0})
    assert d["finance_rate"] == 0.09 and d["ltc"] == 0.55 and d["cost_of_equity"] == 0.14
    assert investor.search_filter(PROFILE, {"postcode": "SW1A 1AA", "strategy": "develop_sell", "asking_price": 3e6})
    assert not investor.search_filter(PROFILE, {"postcode": "M1 1AA", "strategy": "develop_sell"})
    assert not investor.search_filter(PROFILE, {"postcode": "SW1A 1AA", "strategy": "btr"})
    assert not investor.search_filter(PROFILE, {"postcode": "SW1A 1AA", "strategy": "develop_sell", "asking_price": 90e6})


def test_contacts():
    html = '<script type="application/ld+json">{"@type":"RealEstateAgent","name":"Test Agents","telephone":"020 7946 0000","email":"a@b.example"}</script>'
    assert contacts.extract_from_html(html)["phone"] == "020 7946 0000"
    assert contacts.extract_from_html('<a href="tel:+442079460001">call</a><a href="mailto:x@y.example?subject=1">m</a>')["email"] == "x@y.example"
    site = {"name": "Land at X", "source_id": "hamptons", "asking_price": 1.5e6}
    d = contacts.enquiry_draft(site, {"name": "Harry", "firm": "Minerva", "min_deal": 5, "max_deal": 50})
    assert d["to"] == "landandnewhomes@hamptons.co.uk" and "Harry" in d["body"] and "£1,500,000" in d["subject"] + d["body"]


def test_intake_email_and_eml():
    m = outlook.demo_get("demo-1")
    d = intake.heuristic(m["text"], m["subject"], m["from"])
    assert d["postcode"] == "LE11 3TH" and d["asking_price"] == 3_950_000 and d["site_ha"] == 2.9 and d["units"] == 64
    assert d["strategy"] == "develop_sell" and d["planning_status"] == "Consented" and d["bid_deadline"] == "2026-10-23"
    assert sum(x["count"] for x in d["unit_mix"]) == 64 and d["build_cost_total"] == 9_800_000 and d["agent"]["email"].endswith("harcourtvane.example.com")
    eml = (f"From: {m['from']}\nSubject: {m['subject']}\nDate: Fri, 18 Sep 2026 09:42:00 +0000\nContent-Type: text/plain; charset=utf-8\n\n{m['text']}").encode()
    p = intake.parse_eml(eml)
    assert p["subject"].startswith("Off-market") and "Ashby" in p["text"]
    site = intake.to_site(d)
    ap = appraise(site, DEFAULT_ASSUMPTIONS, site["strategy"])
    assert not ap["incomplete"] and ap["units"] == 64


def test_vault_and_analysis():
    s = store.rich_demo_deals()[0]
    assert len(s["docs"]) == 8 and all(docs.path_of(s, d).exists() for d in s["docs"])
    linked = {c["doc"] for c in s["consultees"] if c.get("doc")}
    assert linked <= {d["id"] for d in s["docs"]} and linked
    from core import planning
    res = planning.analyse(docs.all_text(s), s, None)
    assert res["constraints"] and all(c["source_doc"] in {d["id"] for d in s["docs"]} for c in res["constraints"])
    rec = docs.add(s, "Appeal decision APP/X1/W/1.txt", b"Appeal Decision. Planning Inspectorate. Appeal ref APP/X1234/W/25/1")
    assert rec["category"] == "Appeal decision"
    docs.remove(s, rec["id"])
    assert not any(d["id"] == rec["id"] for d in s["docs"])


def test_planning_register():
    deals = {d["id"]: d for d in store.rich_demo_deals()}
    k = planreg.read_out(deals["rich-kingsfield"])
    assert k["status"] == "Consented" and k["permission_expiry"] == "2029-08-13"
    m = planreg.read_out(deals["rich-millpond"])
    assert m["risk"] == "High" and m["open_objections"] == 2
    tl = planreg.timeline(deals["rich-kingsfield"])
    assert tl == sorted(tl, key=lambda e: e["date"]) and len(tl) > 8
    assert any("planning" in l["url"] for l in planreg.links({**deals["rich-kingsfield"], "lat": 53.7, "lon": -1.5}))


def test_comps_parse():
    assert comps.sector("RG40 3QT") == "RG40 3"
    js = {"results": {"bindings": [{"paon": {"value": "12"}, "saon": {"value": "FLAT 3"}, "street": {"value": "MILL LANE"}, "postcode": {"value": "RG40 3QT"}, "amount": {"value": "350000"}, "date": {"value": "2026-03-04"},
                                    "type": {"value": "http://x/flat-maisonette"}, "newbuild": {"value": "true"}, "estate": {"value": "http://x/leasehold"}}]}}
    sales = comps.parse_sparql(js)
    assert comps.attach_areas(sales, {comps._norm("FLAT 3 12 Mill LaneRG40 3QT"): 700.0}) == 1 and sales[0]["psf"] == 500
    assert comps.summarise(sales)["median_psf_flats"] == 500
    assert "error" in comps.fetch({"postcode": ""})


def test_outlook_mocked_graph():
    calls = []

    def post_form(url, fields, timeout=30):
        calls.append(fields.get("grant_type") or "devicecode")
        if "devicecode" in url:
            return {"device_code": "D", "user_code": "ABCD1234", "verification_uri": "https://microsoft.com/devicelogin", "expires_in": 900, "interval": 5}
        if fields["grant_type"].endswith("device_code") and len([c for c in calls if c.endswith("device_code")]) == 1:
            return {"error": "authorization_pending"}
        return {"access_token": "AT", "refresh_token": "RT", "expires_in": 3600}

    def api_get(url, headers=None, timeout=30, cache_ttl=0):
        assert headers["Authorization"] == "Bearer AT"
        if "/attachments" in url:
            return {"value": [{"name": "brochure.txt", "contentBytes": base64.b64encode("Guide price £1m".encode()).decode(), "size": 10}]}
        if "/me/mailFolders" in url:
            return {"value": [{"id": "m1", "subject": "Land for sale, guide price £2m", "from": {"emailAddress": {"name": "A", "address": "a@x.example"}}, "receivedDateTime": "2026-09-01T10:00:00Z", "bodyPreview": "site 2 acres planning", "hasAttachments": True}]}
        if "/me/messages/" in url:
            return {"subject": "S", "from": {"emailAddress": {"name": "A", "address": "a@x.example"}}, "receivedDateTime": "d", "body": {"content": "hello"}}
        return {"mail": "me@corp.example"}
    op, og = outlook.net.post_form, outlook.net.api_get
    outlook.net.post_form, outlook.net.api_get = post_form, api_get
    try:
        outlook.configure("client-1", "common")
        st = outlook.start_device_flow()
        assert st["pending"] and st["device"]["user_code"] == "ABCD1234"
        assert not outlook.poll_device_flow()["connected"]
        st = outlook.poll_device_flow()
        assert st["connected"] and st["account"] == "me@corp.example"
        msgs = outlook.list_messages()
        assert msgs[0]["score"] >= 50
        m = outlook.get_message("m1")
        assert m["attachments"][0][0] == "brochure.txt"
        outlook.disconnect()
        assert not outlook.status()["connected"]
    finally:
        outlook.net.post_form, outlook.net.api_get = op, og
    assert outlook.deal_score("Newsletter unsubscribe", "") == 0 and outlook.deal_score("Lunch?", "") == 0


def test_excel_model_matches_engine():
    """Recalculate every workbook in LibreOffice and compare headline numbers with the Python engine."""
    if not shutil.which("soffice"):
        print("  (skipped: LibreOffice not installed)")
        return
    import openpyxl
    out = pathlib.Path(TMP) / "xl"
    out.mkdir()
    for s in store.rich_demo_deals():
        a = {**DEFAULT_ASSUMPTIONS, **s.get("ovr", {})}
        ap = appraise(s, a, s["strategy"])
        p = out / f"{s['id']}.xlsx"
        xlmodel.build_model(s, a, s["strategy"], ap, None, None, str(p))
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out / "calc"), str(p)], capture_output=True, timeout=180)
        wb = openpyxl.load_workbook(out / "calc" / p.name, data_only=True)
        lab = {}
        for w in ("Appraisal", "Summary"):
            for row in wb[w].iter_rows(values_only=True):
                for i, v in enumerate(row):
                    if isinstance(v, str) and i + 1 < len(row) and isinstance(row[i + 1], (int, float)):
                        lab[v] = row[i + 1]
        for label, key in (("Gross exit value", "gdv"), ("Residual land value (max bid)", "max_bid"), ("Total cost", "total_cost"), ("Profit at land price tested", "profit"), ("Levered IRR", "irr_levered"),
                           ("Unlevered IRR", "irr_unlevered"), ("Equity multiple", "equity_multiple"), ("Peak equity requirement", "peak_equity"), ("NPV of equity at cost of equity", "npv_equity")):
            assert abs(lab[label] - ap[key]) <= max(1, abs(ap[key]) * 1e-4), (s["id"], label, lab[label], ap[key])
        errs = [c.coordinate for w in wb for r in w.iter_rows() for c in r if isinstance(c.value, str) and c.value.startswith(("#", "Err:"))]
        assert not errs, (s["id"], errs[:5])


def test_api_end_to_end():
    import server
    srv = server.make_server(0)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_address[1]}"

    def call(path, body=None):
        r = urllib.request.Request(base + path, data=None if body is None else json.dumps(body).encode(), headers={"Content-Type": "application/json"})
        with urllib.request.urlopen(r) as x:
            raw = x.read()
            return json.loads(raw) if x.headers.get_content_type() == "application/json" else (x.headers.get_content_type(), raw)
    st = call("/api/state")
    ids = [s["id"] for s in st["sites"]]
    assert {"rich-kingsfield", "rich-millpond", "rich-harbour", "rich-infirmary"} <= set(ids)
    call("/api/investor", {"profile": {"active": True, "cost_of_equity": 14, "senior_rate": 9, "regions": ["Yorkshire and the Humber"]}})
    st = call("/api/state")
    assert st["assumptions"]["cost_of_equity"] == 0.14 and st["assumptions"]["finance_rate"] == 0.09
    fit = {s["id"]: s["_ev"]["fit"] for s in st["sites"]}
    assert fit["rich-kingsfield"]["hard_fail"] is False and fit["rich-millpond"]["hard_fail"] is True
    ap = call("/api/appraise", {"id": "rich-kingsfield"})
    assert ap["contact"]["best_phone"] and ap["planning"]["read_out"]["status"] == "Consented" and ap["planning"]["timeline"]
    kind, xl = call("/download/model/rich-harbour.xlsx")
    assert "spreadsheet" in kind and xl[:2] == b"PK"
    doc = next(s for s in st["sites"] if s["id"] == "rich-kingsfield")["docs"][0]
    kind, pdf = call(f"/docs/rich-kingsfield/{doc['id']}")
    assert kind == "application/pdf" and pdf[:4] == b"%PDF"
    r = call("/api/outlook/import", {"mid": "demo-1", "demo": True})
    site = next(s for s in call("/api/state")["sites"] if s["id"] == r["id"])
    assert site["_ev"]["appraisal"]["units"] == 64 and site["_ndocs"] == 1 and site["_contact"]["best_email"]
    saved = call("/api/model/save", {"id": r["id"]})
    assert pathlib.Path(saved["path"]).exists()
    call("/api/planning/entry", {"id": "rich-millpond", "kind": "consultee", "op": "save", "index": None, "entry": {"body": "Police", "stance": "No objection", "summary": "Secured by Design", "date": "2026-09-01"}})
    assert any(c["body"] == "Police" for c in call("/api/planning/get", {"id": "rich-millpond"})["consultees"])
    up = call("/api/docs/upload", {"id": "rich-millpond", "files": [{"name": "Committee report.txt", "data": base64.b64encode(b"Officer report to committee. Case officer recommends approval.").decode()}]})
    assert up["added"][0]["category"] == "Officer report"
    assert call("/api/comps", {"id": "rich-kingsfield"})["demo"] is True
    assert call("/api/comps", {"id": r["id"]}).get("error")  # offline: graceful message, no crash


def test_hosted_password_wall():
    import http.client
    import server
    server.PASSWORD, server.SECRET, server.HOSTED = "s3cret", b"s3cret", True
    try:
        srv = server.make_server(0)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        port = srv.server_address[1]

        def req(method, path, body=None, cookie=None, headers=None):
            c = http.client.HTTPConnection("127.0.0.1", port)
            h = {"Cookie": cookie} if cookie else {}
            h.update(headers or {})
            c.request(method, path, body=body, headers=h)
            r = c.getresponse()
            return r.status, dict(r.getheaders()), r.read()
        assert req("GET", "/healthz")[0] == 200
        st, h, _ = req("GET", "/")
        assert st == 303 and h["Location"] == "/login"
        assert req("GET", "/api/state")[0] == 401
        assert req("GET", "/docs/rich-kingsfield/x")[0] == 303
        assert req("POST", "/login", "password=wrong", headers={"Content-Type": "application/x-www-form-urlencoded", "Content-Length": "14"})[0] == 401
        st, h, _ = req("POST", "/login", "password=s3cret", headers={"Content-Type": "application/x-www-form-urlencoded", "Content-Length": "15", "X-Forwarded-Proto": "https"})
        assert st == 303 and "HttpOnly" in h["Set-Cookie"] and "Secure" in h["Set-Cookie"]
        ck = h["Set-Cookie"].split(";")[0]
        st, _, body = req("GET", "/api/state", cookie=ck)
        assert st == 200 and json.loads(body)["hosted"] is True
        assert req("GET", "/api/state", cookie="minerva_session=1.deadbeef")[0] == 401
        assert req("GET", "/download/model/rich-harbour.xlsx", cookie=ck)[0] == 200
    finally:
        server.PASSWORD, server.SECRET, server.HOSTED = "", b"local", False


if __name__ == "__main__":
    for n, f in list(globals().items()):
        if n.startswith("test_"):
            f()
            print("ok", n)
