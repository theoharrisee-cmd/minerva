"""Excel financial model export: a live, formula-driven workbook (Inputs > Appraisal > Cash Flow > Summary).
Sensitivity and scenarios are computed by the Minerva engine and pasted as values (they refresh on re-export)."""
from __future__ import annotations

from copy import deepcopy
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from .appraisal import appraise
from .config import ASSUMPTION_SPEC, DEFAULT_ASSUMPTIONS, STRATEGIES

NAVY, TEAL, GREY = "0B3F62", "0F766E", "F3F5F7"
BLUE = Font(name="Arial", size=10, color="0000FF")
BLK = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FILL_H = PatternFill("solid", fgColor=NAVY)
FILL_S = PatternFill("solid", fgColor=GREY)
FILL_IN = PatternFill("solid", fgColor="FFF9E5")
thin = Side(style="thin", color="D5DAE0")
GBP, PCT, NUM = '£#,##0;[Red](£#,##0);"-"', "0.0%", "#,##0;[Red](#,##0);\"-\""


def _sheet(wb, title, widths):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w
    return ws


def _title(ws, text, sub=""):
    ws["B2"] = text
    ws["B2"].font = Font(name="Arial", size=16, bold=True, color=NAVY)
    if sub:
        ws["B3"] = sub
        ws["B3"].font = Font(name="Arial", size=9, italic=True, color="6B7686")


def _band(ws, row, text, cols=5):
    for c in range(2, 2 + cols):
        ws.cell(row, c).fill = FILL_H
    ws.cell(row, 2, text).font = HDR


def build_model(site: dict, assumptions: dict, strategy: str, ap: dict, investor: dict | None, fit: dict | None, path: str) -> dict:
    a = {**DEFAULT_ASSUMPTIONS, **assumptions}
    strat = strategy
    wb = Workbook()
    wb.remove(wb.active)
    R: dict[str, str] = {}

    # ------------------------------------------------------------------ Inputs
    wi = _sheet(wb, "Inputs", [3, 46, 16, 14, 14, 14, 44])
    _title(wi, f"{site.get('name', 'Deal')}: inputs", f"Blue cells are inputs. Strategy: {STRATEGIES[strat]}. Generated {date.today():%d %b %Y} by Minerva.")
    r = 5

    def band(text):
        nonlocal r
        _band(wi, r, text, 6)
        r += 1

    def inp(key, label, val, fmt=NUM, note="", src=None):
        nonlocal r
        wi.cell(r, 2, label).font = BLK
        c = wi.cell(r, 3, val)
        c.font = BLUE
        c.fill = FILL_IN
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
        if note:
            wi.cell(r, 7, note).font = Font(name="Arial", size=9, italic=True, color="6B7686")
        R[key] = f"Inputs!$C${r}"
        r += 1

    band("Deal")
    wi.cell(r, 2, "Name").font = BLK
    wi.cell(r, 3, site.get("name", "")).font = BLUE
    r += 1
    wi.cell(r, 2, "Address").font = BLK
    wi.cell(r, 3, site.get("address", "")).font = BLUE
    r += 1
    inp("asking", "Asking price (£)", float(site.get("asking_price") or 0), GBP)
    inp("land", "Land price tested (£)", float(ap["tested_land_price"]), GBP, "Defaults to the asking price. Overwrite to test a bid.")
    band("Scheme")
    inp("eff", "Net to gross efficiency", a["efficiency"], PCT)
    inp("gia_over", "GIA override (sq ft, 0 = derive)", float(ap["gia"]) if (site.get("gia_sqft") or strat == "value_add") else 0.0, NUM)
    inp("aff", "Affordable housing share", a["affordable_pct"], PCT)
    inp("aff_disc", "Affordable discount to market", a["affordable_disc"], PCT)
    band("Costs")
    inp("build_psf", "Build / refurbishment (£ per sq ft GIA)", a["refurb_psf"] if strat == "value_add" else a["build_psf"], GBP)
    inp("build_tot", "Build cost total override (£, 0 = derive)", float(site.get("build_cost_total") or 0), GBP, "Quantity surveyor cost plan, if you have one")
    inp("ext", "Externals and abnormals (% of build)", 0.0 if strat == "value_add" else a["external_pct"], PCT)
    inp("cont", "Contingency (% of build)", a["contingency_pct"], PCT)
    inp("fees", "Professional fees (% of build)", a["fees_pct"], PCT)
    inp("s106", "S106 per unit (£)", 0.0 if strat == "value_add" else a["s106_per_unit"], GBP)
    inp("cil", "CIL (£ per sq ft GIA)", 0.0 if strat == "value_add" else a["cil_psf"], GBP)
    inp("mkt", "Marketing and agents (% of GDV)", a["marketing_pct"], PCT)
    inp("legal", "Sales legal (% of GDV)", a["sales_legal_pct"], PCT)
    inp("acq", "Land acquisition costs (%)", a["acq_costs_pct"], PCT)
    band("Income and exit")
    inp("void", "Void and bad debt", a["void_pct"], PCT)
    inp("opex", "Operating costs (% of gross rent)", a["opex_pct"], PCT)
    inp("yield", "Exit yield (net initial)", a["exit_yield"], "0.00%")
    inp("pcost", "Purchaser costs on exit", a["purchaser_costs"], PCT)
    band("Programme")
    inp("pl", "Months to secure planning", a["planning_months"], NUM)
    inp("bm", "Build months", ap["build_months"], NUM)
    inp("sm", "Sales / lease-up months", a["sales_months"], NUM)
    band("Finance and returns")
    inp("rate", "Senior debt rate (pa)", a["finance_rate"], "0.00%")
    inp("ltc", "Senior loan to cost", a["ltc"], PCT)
    inp("mrate", "Mezzanine rate (pa)", a["mezz_rate"], "0.00%")
    inp("mltc", "Mezzanine loan to cost (additional)", a["mezz_ltc"], PCT)
    inp("fee", "Debt arrangement fee (%)", a["debt_fee_pct"], "0.00%")
    inp("coe", "Cost of equity / hurdle rate (pa)", a["cost_of_equity"], PCT)
    p_key = {"develop_sell": "target_profit_gdv", "btr": "target_profit_cost_btr", "value_add": "target_profit_cost_va"}[strat]
    inp("target", f"Target profit on {'GDV' if strat == 'develop_sell' else 'cost'}", a[p_key], PCT)

    r += 1
    _band(wi, r, "Unit mix (edit rows freely; keep within the table)", 6)
    r += 1
    for i, h in enumerate(["Type", "Units", "NSA sq ft", "Sale value £ / unit", "Rent £ pcm / unit"]):
        c = wi.cell(r, 2 + i, h)
        c.font = BOLD
        c.fill = FILL_S
    r += 1
    mix = [m for m in (site.get("unit_mix") or []) if m.get("count")]
    if strat == "value_add" or not mix:
        units = ap["units"]
        nsa = ap["nsa"]
        if strat == "value_add":
            sale_total = nsa * ap["sales_psf"] * (1 + a["uplift_pct"])
            rent_total = nsa * ap["rent_psf"] * (1 + a["uplift_pct"])
        else:
            sale_total = nsa * ap["sales_psf"]
            rent_total = nsa * ap["rent_psf"]
        mix = [dict(type="Average unit" if strat != "value_add" else "Whole building", count=units, nsa_sqft=nsa / units, sale_value=sale_total / units, rent_pcm=rent_total / 12 / units)]
    first = r
    for m in mix:
        nsa_u = float(m.get("nsa_sqft") or a["unit_nsa_sqft"])
        vals = [m.get("type", ""), m["count"], nsa_u, float(m.get("sale_value") or nsa_u * ap["sales_psf"]), float(m.get("rent_pcm") or nsa_u * ap["rent_psf"] / 12)]
        for i, v in enumerate(vals):
            c = wi.cell(r, 2 + i, v)
            c.font = BLUE
            c.fill = FILL_IN
            if i:
                c.number_format = NUM if i < 3 else GBP
        r += 1
    last = r - 1
    for _ in range(3):  # spare rows
        for i in range(5):
            c = wi.cell(r, 2 + i)
            c.fill = FILL_IN
            c.font = BLUE
            if i:
                c.number_format = NUM if i < 3 else GBP
        last = r
        r += 1
    rng = lambda col: f"Inputs!${col}${first}:${col}${last}"
    wi.freeze_panes = "A5"

    # --------------------------------------------------------------- Appraisal
    wa = _sheet(wb, "Appraisal", [3, 46, 18, 3, 60])
    _title(wa, "Appraisal", "All figures are formulas driven by the Inputs sheet.")
    A: dict[str, str] = {}
    row = [5]

    def sec(text):
        _band(wa, row[0], text, 4)
        row[0] += 1

    def ln(key, label, formula, fmt=GBP, bold=False, note=""):
        wa.cell(row[0], 2, label).font = BOLD if bold else BLK
        c = wa.cell(row[0], 3, formula)
        c.number_format = fmt
        c.font = BOLD if bold else BLK
        if note:
            wa.cell(row[0], 5, note).font = Font(name="Arial", size=9, italic=True, color="6B7686")
        A[key] = f"Appraisal!$C${row[0]}"
        row[0] += 1

    sec("Scheme")
    ln("units", "Units", f"=SUM({rng('C')})", NUM)
    ln("nsa", "Net sellable area (sq ft)", f"=SUMPRODUCT({rng('C')},{rng('D')})", NUM)
    ln("gia", "Gross internal area (sq ft)", f"=IF({R['gia_over']}>0,{R['gia_over']},{A['nsa']}/{R['eff']})", NUM)
    ln("afac", "Affordable value factor", "=1" if strat == "value_add" else f"=(1-{R['aff']})+{R['aff']}*(1-{R['aff_disc']})", "0.000")
    sec("Costs excluding land")
    ln("build", "Build / refurbishment", f"=IF({R['build_tot']}>0,{R['build_tot']},{A['gia']}*{R['build_psf']}*(1+{R['ext']}))")
    ln("cont", "Contingency", f"={A['build']}*{R['cont']}")
    ln("fees", "Professional fees", f"={A['build']}*{R['fees']}")
    ln("s106", "S106", f"={A['units']}*{R['s106']}")
    ln("cil", "CIL", f"={A['gia']}*{R['cil']}")
    ln("cpre", "Total costs before sales costs and finance", f"=SUM({A['build']}:{A['cil']})".replace("Appraisal!", ""), bold=True)
    sec("Revenue and value")
    ln("mkt_gdv", "Market value of units", f"=SUMPRODUCT({rng('C')},{rng('E')})")
    ln("gdv", "GDV after affordable discount", f"={A['mkt_gdv']}*{A['afac']}")
    ln("rent", "Gross annual rent", f"=SUMPRODUCT({rng('C')},{rng('F')})*12*{A['afac']}")
    ln("noi", "Net operating income", f"={A['rent']}*(1-{R['void']})*(1-{R['opex']})")
    ln("ivalue", "Income value (NOI / exit yield)", f"=IF({R['yield']}>0,{A['noi']}/{R['yield']},0)")
    gv = {"develop_sell": f"={A['gdv']}", "btr": f"={A['ivalue']}", "value_add": f"=MAX({A['ivalue']},{A['gdv']})"}[strat]
    ln("gross", "Gross exit value", gv, bold=True, note={"develop_sell": "Sell: GDV", "btr": "BTR: capitalised income", "value_add": "Value-add: higher of resale and income value"}[strat])
    sc = f"={A['gdv']}*({R['mkt']}+{R['legal']})" if strat == "develop_sell" else f"={A['gross']}*{R['pcost']}/(1+{R['pcost']})"
    ln("scost", "Sales / purchaser costs", sc)
    ln("net", "Net exit value", f"={A['gross']}-{A['scost']}", bold=True)
    sec("Programme and finance")
    ln("pl", "Planning months (rounded)", f"=ROUND({R['pl']},0)", NUM)
    ln("bm", "Build months (rounded)", f"=ROUND({R['bm']},0)", NUM)
    ln("sm", "Sales months (rounded)", f"=MAX(1,ROUND({R['sm']},0))", NUM)
    ln("lm", "Months land is financed", f"={R['pl']}+{R['bm']}+{R['sm']}", NUM)
    ln("k", "Land cost factor (costs + finance)", f"=1+{R['acq']}+{R['rate']}*{A['lm']}/12", "0.0000")
    ln("bfin", "Build finance", f"={R['rate']}*({A['cpre']}*0.55*{R['bm']}/12+{A['cpre']}*{R['sm']}/2/12)")
    sec("Residual land value")
    if strat == "develop_sell":
        f = f"=MAX(0,({A['net']}-{A['cpre']}-{A['bfin']}-{R['target']}*{A['gdv']})/{A['k']})"
    else:
        f = f"=MAX(0,({A['net']}/(1+{R['target']})-{A['cpre']}-{A['bfin']})/{A['k']})"
    ln("rlv", "Residual land value (max bid)", f, bold=True, note="Land price at which the target profit is exactly met")
    ln("offer", "Opening offer", f"={A['rlv']}*(1-{a['offer_discount']})", note=f"{a['offer_discount']*100:.0f}% below max bid")
    ln("head", "Headroom vs asking", f"=IF({R['asking']}>0,({A['rlv']}-{R['asking']})/{R['asking']},\"n/a\")", PCT)
    sec("Appraisal at the land price tested")
    ln("lcost", "Land incl. acquisition costs", f"={R['land']}*(1+{R['acq']})")
    ln("lfin", "Land finance", f"={R['land']}*{R['rate']}*{A['lm']}/12")
    ln("tcost", "Total cost", f"={A['lcost']}+{A['lfin']}+{A['cpre']}" + (f"+{A['scost']}" if strat == "develop_sell" else "") + f"+{A['bfin']}", bold=True)
    ln("profit", "Profit", f"={A['gross']}-{A['tcost']}" if strat == "develop_sell" else f"={A['net']}-{A['tcost']}", bold=True)
    ln("pgdv", "Profit on GDV", f"=IF({A['gross']}>0,{A['profit']}/{A['gross']},0)", PCT)
    ln("pcost_", "Profit on cost", f"=IF({A['tcost']}>0,{A['profit']}/{A['tcost']},0)", PCT)
    ln("yoc", "Yield on cost", f"=IF({A['tcost']}>0,{A['noi']}/{A['tcost']},0)", "0.00%")
    basis = A["pgdv"] if strat == "develop_sell" else A["pcost_"]
    ln("viable", "Meets target profit?", f"=IF({basis}>={R['target']}-0.000000001,\"Yes\",\"No\")", "@", bold=True)

    # --------------------------------------------------------------- Cash flow
    N = max(72, int(round(a["planning_months"] + ap["build_months"] + a["sales_months"])) + 12)
    wc = _sheet(wb, "Cash Flow", [3, 34] + [11] * (N + 1))
    _title(wc, "Monthly cash flow and funding", "Month 0 is land purchase. Build spend follows an S-curve. Senior debt and mezzanine draw pro rata to costs; interest rolls up.")
    cols = list(range(3, 3 + N + 1))
    labels = ["Month", "Land and acquisition costs", "Build and other costs", "Net exit proceeds", "Unlevered net cash flow", "Cumulative unlevered",
              "Debt fee", "Net after fee",
              "Senior: opening", "Senior: interest", "Senior: draw", "Senior: repayment", "Senior: closing (pre-exit)", "Senior: closing",
              "Mezzanine: opening", "Mezzanine: interest", "Mezzanine: draw", "Mezzanine: repayment", "Mezzanine: closing (pre-exit)", "Mezzanine: closing",
              "Equity cash flow", "Cumulative equity"]
    base = 5
    rowi = {lab: base + i for i, lab in enumerate(labels)}
    for lab, ri in rowi.items():
        wc.cell(ri, 2, lab).font = BOLD if lab in ("Month", "Unlevered net cash flow", "Equity cash flow") else BLK
    for c in cols:
        col = L(c)
        prev = L(c - 1)
        m = f"{col}${rowi['Month']}"
        first_col = c == cols[0]
        wc.cell(rowi["Month"], c, c - 3).font = HDR
        wc.cell(rowi["Month"], c).fill = FILL_H
        S = lambda x: f"(3*({x})^2-2*({x})^3)"
        wc.cell(rowi["Land and acquisition costs"], c, f"=-IF({m}=0,{R['land']}*(1+{R['acq']}),0)")
        x1 = f"MIN(1,MAX(0,({m}-{A['pl']})/{A['bm']}))"
        x0 = f"MIN(1,MAX(0,({m}-{A['pl']}-1)/{A['bm']}))"
        wc.cell(rowi["Build and other costs"], c, f"=-IF(AND({m}>={A['pl']}+1,{m}<={A['pl']}+{A['bm']}),{A['cpre']}*({S(x1)}-{S(x0)}),0)")
        if strat == "develop_sell":
            wc.cell(rowi["Net exit proceeds"], c, f"=IF(AND({m}>={A['pl']}+{A['bm']}+1,{m}<={A['pl']}+{A['bm']}+{A['sm']}),{A['net']}/{A['sm']},0)")
        else:
            wc.cell(rowi["Net exit proceeds"], c, f"=IF({m}={A['pl']}+{A['bm']}+{A['sm']},{A['net']},0)")
        u = f"{col}{rowi['Unlevered net cash flow']}"
        wc.cell(rowi["Unlevered net cash flow"], c, f"=SUM({col}{rowi['Land and acquisition costs']}:{col}{rowi['Net exit proceeds']})")
        wc.cell(rowi["Cumulative unlevered"], c, f"={u}" if first_col else f"={prev}{rowi['Cumulative unlevered']}+{u}")
        wc.cell(rowi["Debt fee"], c, f"=IF({m}=0,{R['fee']}*({R['ltc']}+{R['mltc']})*-SUMIF($C${rowi['Unlevered net cash flow']}:${L(cols[-1])}${rowi['Unlevered net cash flow']},\"<0\"),0)")
        cf = f"{col}{rowi['Net after fee']}"
        wc.cell(rowi["Net after fee"], c, f"={u}-{col}{rowi['Debt fee']}")
        lastm = A["lm"]
        for t, (nm, rt, lt) in {"Senior": ("Senior", R["rate"], R["ltc"]), "Mezzanine": ("Mezzanine", R["mrate"], R["mltc"])}.items():
            o, i_, d, rp, pc, cl = (f"{nm}: {x}" for x in ("opening", "interest", "draw", "repayment", "closing (pre-exit)", "closing"))
            wc.cell(rowi[o], c, 0 if first_col else f"={prev}{rowi[cl]}")
            wc.cell(rowi[i_], c, f"={col}{rowi[o]}*{rt}/12")
            wc.cell(rowi[d], c, f"=MAX(0,-{cf})*{lt}")
            if t == "Senior":
                wc.cell(rowi[rp], c, f"=MIN({col}{rowi[o]}+{col}{rowi[i_]},MAX({cf},0))")
            else:
                wc.cell(rowi[rp], c, f"=MIN({col}{rowi[o]}+{col}{rowi[i_]},MAX({cf},0)-{col}{rowi['Senior: repayment']})")
            wc.cell(rowi[pc], c, f"={col}{rowi[o]}+{col}{rowi[i_]}+{col}{rowi[d]}-{col}{rowi[rp]}")
            wc.cell(rowi[cl], c, f"=IF({m}>={lastm},0,{col}{rowi[pc]})")
        wc.cell(rowi["Equity cash flow"], c, f"={cf}+{col}{rowi['Senior: draw']}+{col}{rowi['Mezzanine: draw']}-{col}{rowi['Senior: repayment']}-{col}{rowi['Mezzanine: repayment']}"
                f"-IF({m}={lastm},{col}{rowi['Senior: closing (pre-exit)']}+{col}{rowi['Mezzanine: closing (pre-exit)']},0)")
        e = f"{col}{rowi['Equity cash flow']}"
        wc.cell(rowi["Cumulative equity"], c, f"={e}" if first_col else f"={prev}{rowi['Cumulative equity']}+{e}")
        for lab, ri in rowi.items():
            if lab != "Month":
                wc.cell(ri, c).number_format = NUM
                wc.cell(ri, c).font = BLK
    wc.freeze_panes = "C6"
    e_rng = f"'Cash Flow'!$C${rowi['Equity cash flow']}:${L(cols[-1])}${rowi['Equity cash flow']}"
    u_rng = f"'Cash Flow'!$C${rowi['Unlevered net cash flow']}:${L(cols[-1])}${rowi['Unlevered net cash flow']}"
    cum_rng = f"'Cash Flow'!$C${rowi['Cumulative equity']}:${L(cols[-1])}${rowi['Cumulative equity']}"
    month_rng = f"'Cash Flow'!$C${rowi['Month']}:${L(cols[-1])}${rowi['Month']}"

    # ----------------------------------------------------------------- Summary
    ws = _sheet(wb, "Summary", [3, 40, 22, 3, 40, 22])
    _title(ws, site.get("name", "Deal"), f"{STRATEGIES[strat]} · {site.get('address', '')}")
    _band(ws, 5, "Headline returns", 5)
    rows = [
        ("Gross exit value", f"={A['gross']}", GBP), ("Total cost", f"={A['tcost']}", GBP), ("Profit at land price tested", f"={A['profit']}", GBP),
        ("Profit on GDV", f"={A['pgdv']}", PCT), ("Profit on cost", f"={A['pcost_']}", PCT), ("Target profit", f"={R['target']}", PCT), ("Meets target?", f"={A['viable']}", "@"),
        ("Residual land value (max bid)", f"={A['rlv']}", GBP), ("Opening offer", f"={A['offer']}", GBP), ("Asking price", f"={R['asking']}", GBP), ("Headroom vs asking", f"={A['head']}", PCT),
    ]
    rows2 = [
        ("Unlevered IRR", f"=IFERROR((1+IRR({u_rng},0.01))^12-1,\"n/a\")", PCT),
        ("Levered IRR", f"=IFERROR((1+IRR({e_rng},0.01))^12-1,\"n/a\")", PCT),
        ("Equity multiple", f"=IFERROR(SUMIF({e_rng},\">0\")/-SUMIF({e_rng},\"<0\"),0)", "0.00\"x\""),
        ("Peak equity requirement", f"=-MIN({cum_rng})", GBP),
        ("NPV of equity at cost of equity", f"='Cash Flow'!C{rowi['Equity cash flow']}+NPV((1+{R['coe']})^(1/12)-1,'Cash Flow'!D{rowi['Equity cash flow']}:{L(cols[-1])}{rowi['Equity cash flow']})", GBP),
        ("Cost of equity (hurdle)", f"={R['coe']}", PCT), ("Units", f"={A['units']}", NUM), ("Net sellable area (sq ft)", f"={A['nsa']}", NUM),
        ("Yield on cost", f"={A['yoc']}", "0.00%"), ("Net operating income", f"={A['noi']}", GBP), ("Months land financed", f"={A['lm']}", NUM),
    ]
    for i, (lab, f, fmt) in enumerate(rows):
        ws.cell(6 + i, 2, lab).font = BLK
        c = ws.cell(6 + i, 3, f)
        c.number_format = fmt
        c.font = BOLD
        c.alignment = Alignment(horizontal="right")
    for i, (lab, f, fmt) in enumerate(rows2):
        ws.cell(6 + i, 5, lab).font = BLK
        c = ws.cell(6 + i, 6, f)
        c.number_format = fmt
        c.font = BOLD
        c.alignment = Alignment(horizontal="right")
    ws["E5"].fill = FILL_H
    ws["F5"].fill = FILL_H
    ws["E5"] = "Returns to equity"
    ws["E5"].font = HDR

    rr = 19
    if investor and investor.get("active"):
        _band(ws, rr, f"Mandate fit: {fit.get('label', '')}" if fit else "Mandate", 5)
        rr += 1
        for c in (fit or {}).get("checks", []):
            ws.cell(rr, 2, c["name"]).font = BLK
            ws.cell(rr, 3, "Pass" if c["ok"] else "Fail" if c["ok"] is False else "n/a").font = BOLD
            ws.cell(rr, 5, c["detail"]).font = Font(name="Arial", size=9, color="6B7686")
            rr += 1
        rr += 1

    # Scenarios (engine values)
    _band(ws, rr, "Scenarios (values from the Minerva engine at the land price tested)", 5)
    rr += 1
    for i, h in enumerate(["Scenario", "Profit", "Profit on cost", "Levered IRR", "Equity multiple"]):
        c = ws.cell(rr, [2, 3, 5, 6, 7][i], h)
        c.font = BOLD
        c.fill = FILL_S
    ws.column_dimensions["G"].width = 16
    rr += 1
    land = ap["tested_land_price"]
    base_site = dict(site)
    for name, ds, dc, dm in [("Downside: sales -10%, build +10%, +6 months", -0.10, 0.10, 6), ("Base case", 0, 0, 0), ("Upside: sales +5%, build -5%", 0.05, -0.05, 0)]:
        a2, s2 = deepcopy(a), dict(base_site)
        s2["sales_psf"] = (site.get("sales_psf") or a["sales_psf"]) * (1 + ds)
        s2["rent_psf"] = (site.get("rent_psf") or a["rent_psf"]) * (1 + ds)
        a2["build_psf"] *= 1 + dc
        a2["refurb_psf"] *= 1 + dc
        if s2.get("build_cost_total"):
            s2["build_cost_total"] *= 1 + dc
        a2["build_months"] += dm
        try:
            x = appraise(s2, a2, strat, land_price=land)
            vals = [x["profit"], x["profit_on_cost"], x["irr_levered"], x["equity_multiple"]]
        except Exception:
            vals = [None] * 4
        ws.cell(rr, 2, name).font = BLK
        for j, (v, fmt) in enumerate(zip(vals, [GBP, PCT, PCT, '0.00"x"'])):
            c = ws.cell(rr, [3, 5, 6, 7][j], v if v is not None else "n/a")
            c.number_format = fmt
            c.font = BLK
        rr += 1

    # -------------------------------------------------------------- Sensitivity
    wsn = _sheet(wb, "Sensitivity", [3, 24] + [12] * 7)
    _title(wsn, "Sensitivity", "Values from the Minerva engine at export time. Re-export after changing inputs.")
    from .appraisal import sensitivity
    steps = [-0.15, -0.10, -0.05, 0, 0.05, 0.10, 0.15]
    ykey = "sales_psf" if strat in ("develop_sell", "value_add") else "exit_yield"
    xkey = "build_psf" if strat != "value_add" else "refurb_psf"
    metric = "profit_on_gdv" if strat == "develop_sell" else "profit_on_cost"
    sens = sensitivity(site, a, xkey, steps, ykey, steps, metric, strat)
    wsn["B5"] = f"Profit on {'GDV' if strat == 'develop_sell' else 'cost'}. Rows: {ASSUMPTION_SPEC[ykey][1]}. Columns: {ASSUMPTION_SPEC[xkey][1]}."
    wsn["B5"].font = BOLD
    for j, s_ in enumerate(steps):
        c = wsn.cell(7, 3 + j, s_)
        c.number_format = "+0%;-0%;0%"
        c.font = HDR
        c.fill = FILL_H
        c = wsn.cell(8 + j, 2, s_)
        c.number_format = "+0%;-0%;0%"
        c.font = HDR
        c.fill = FILL_H
        for i, v in enumerate(sens["z"][j]):
            cc = wsn.cell(8 + j, 3 + i, v)
            cc.number_format = PCT
            cc.font = BLK
            cc.fill = PatternFill("solid", fgColor="D8F0E0" if v >= a[p_key] else "FBE3DE")

    # --------------------------------------------------------------------- Info
    wn = _sheet(wb, "Deal Info", [3, 26, 100])
    _title(wn, "Deal information")
    info = [("Description", site.get("description")), ("Tenure", site.get("tenure")), ("Vendor", site.get("vendor")), ("Marketing", site.get("marketing")),
            ("Bid deadline", site.get("bid_deadline")), ("Planning status", site.get("planning_status")), ("Notes", site.get("notes"))]
    ag = site.get("agent") or {}
    info += [("Agent", " · ".join(x for x in (ag.get("name"), ag.get("firm"), ag.get("phone"), ag.get("email")) if x))]
    for h in site.get("planning_history", []):
        info.append((f"Planning {h['ref']}", f"{h['status']}: {h['description']}"))
    for c in site.get("consultees", []):
        info.append((f"Consultee: {c['body']}", f"{c['stance']}: {c['summary']}"))
    for d in site.get("docs", []):
        info.append(("Document", f"{d['name']} ({d['category']})"))
    for i, (k, v) in enumerate(info):
        if v:
            wn.cell(5 + i, 2, k).font = BOLD
            c = wn.cell(5 + i, 3, v)
            c.font = BLK
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.move_sheet("Summary", offset=-4)
    wb.move_sheet("Inputs", offset=-3)
    wb.active = 0
    wb.save(path)
    return {"summary": {k: v for k, v in A.items()}, "cf_rows": rowi}
